from io import BytesIO

from docx import Document
from reportlab.pdfgen import canvas

from autoinsight.loaders import load_bytes


def test_markdown_and_csv_locations():
    chunks = load_bytes("sample.md", "# 充电\n30% 至 80% 需要 28 分钟。".encode())
    assert len(chunks) == 1
    assert chunks[0].citation == "sample.md · 充电"
    rows = load_bytes("voc.csv", "主题,反馈\n充电,状态不准\n语音,识别慢\n".encode())
    assert [row.location for row in rows] == ["行 2", "行 3"]


def test_pdf_page_and_docx_heading_locations():
    pdf = BytesIO()
    page = canvas.Canvas(pdf)
    page.drawString(80, 750, "Charging takes 28 minutes.")
    page.showPage()
    page.save()
    pdf_chunks = load_bytes("guide.pdf", pdf.getvalue())
    assert pdf_chunks[0].citation == "guide.pdf · 第 1 页"
    assert "28 minutes" in pdf_chunks[0].text

    doc = Document()
    doc.add_heading("售后流程", 1)
    doc.add_paragraph("三天内联系用户。")
    doc_file = BytesIO()
    doc.save(doc_file)
    word_chunks = load_bytes("guide.docx", doc_file.getvalue())
    assert "售后流程" in word_chunks[0].location
    assert "三天内" in word_chunks[0].text


def test_decimal_is_not_split():
    chunks = load_bytes("screen.md", "# 屏幕\n尺寸是 15.6 英寸。".encode())
    assert "15.6 英寸" in chunks[0].text
