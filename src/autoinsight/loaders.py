from __future__ import annotations

import csv
import hashlib
import io
import re
from pathlib import Path

from .models import Chunk

SUPPORTED = {".md", ".txt", ".csv", ".pdf", ".docx"}


def load_paths(paths: list[Path]) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in sorted(paths, key=lambda item: str(item).lower()):
        if path.is_dir():
            children = [p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED]
            chunks.extend(load_paths(children))
        elif path.suffix.lower() in SUPPORTED:
            chunks.extend(load_bytes(path.name, path.read_bytes()))
    return chunks


def load_bytes(name: str, data: bytes) -> list[Chunk]:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}")
    source = Path(name).name
    if suffix in {".md", ".txt"}:
        text = data.decode("utf-8-sig")
        sections = _text_sections(text, markdown=suffix == ".md")
    elif suffix == ".csv":
        rows = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        sections = [(f"行 {index}", "；".join(f"{key}: {value}" for key, value in row.items() if value))
                    for index, row in enumerate(rows, 2)]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        sections = [(f"第 {index} 页", page.extract_text() or "")
                    for index, page in enumerate(reader.pages, 1)]
    else:
        from docx import Document

        doc = Document(io.BytesIO(data))
        sections = _docx_sections(doc)
    output: list[Chunk] = []
    for location, section in sections:
        for part_index, piece in enumerate(_split(section), 1):
            suffix_label = f" · 片段 {part_index}" if len(section) > 500 else ""
            identity = hashlib.sha256(f"{source}|{location}|{part_index}|{piece}".encode()).hexdigest()[:12]
            output.append(Chunk(identity, source, location + suffix_label, piece))
    return output


def _text_sections(text: str, markdown: bool) -> list[tuple[str, str]]:
    title = "正文"
    sections: list[tuple[str, str]] = []
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if buffer:
            sections.append((title, "\n".join(buffer)))
            buffer = []

    for line in text.splitlines():
        if markdown and re.match(r"^#{1,6}\s+", line):
            flush()
            title = re.sub(r"^#{1,6}\s+", "", line).strip()
        elif line.strip():
            buffer.append(line.strip())
        else:
            flush()
    flush()
    return sections


def _docx_sections(doc: object) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    heading = "正文"
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            heading = text
        else:
            sections.append((f"章节 {heading} · 段落 {len(sections) + 1}", text))
    for table_index, table in enumerate(doc.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                sections.append((f"表 {table_index} 行 {row_index}", "；".join(cells)))
    return sections


def _split(text: str, limit: int = 500) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    paragraphs = re.split(r"(?<=[。！？；!?])\s*|(?<!\d)\.\s+|\n", text)
    pieces: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if not paragraph:
            continue
        if len(current) + len(paragraph) > limit and current:
            pieces.append(current.strip())
            current = ""
        while len(paragraph) > limit:
            pieces.append(paragraph[:limit].strip())
            paragraph = paragraph[limit:]
        current += paragraph + " "
    if current.strip():
        pieces.append(current.strip())
    return pieces
