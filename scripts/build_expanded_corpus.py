"""Deterministically build the redistributable expanded offline evaluation corpus."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "expanded"
CASES = ROOT / "data" / "evaluation" / "expanded_questions.jsonl"

# These are original Chinese summaries of small factual points, not copied pages.
PUBLIC = [
    ("纯电动车通过外接电源给动力电池充电。", "纯电动车的动力电池如何补能？", "外接电源", "https://afdc.energy.gov/vehicles/electric-basics-ev"),
    ("纯电动车没有本地尾气排放；发电环节仍可能产生排放。", "纯电动车是否有本地尾气排放？", "没有本地尾气排放", "https://afdc.energy.gov/vehicles/electric-basics-ev"),
    ("低温会增加座舱加热能耗，因此实际续航可能下降。", "低温为何可能降低电车续航？", "座舱加热能耗", "https://afdc.energy.gov/vehicles/electric-basics-ev"),
    ("频繁制动可回收部分能量，纯电动车在城市工况往往更高效。", "为什么纯电动车在城市工况可能更高效？", "回收部分能量", "https://afdc.energy.gov/vehicles/electric-basics-ev"),
    ("美国交流一级充电使用 120 V 电源。", "美国交流一级充电使用多少伏电源？", "120 V", "https://afdc.energy.gov/fuels/electricity-stations"),
    ("美国住宅交流二级充电通常使用 240 V 电源。", "美国住宅交流二级充电通常使用多少伏？", "240 V", "https://afdc.energy.gov/fuels/electricity-stations"),
    ("充电时间受电池剩余电量、容量、车载充电机和充电设备功率影响。", "哪些因素会影响充电时间？", "电池剩余电量", "https://afdc.energy.gov/fuels/electricity-stations"),
    ("插电式混合动力汽车既能外接充电，也配有内燃机。", "插电式混合动力汽车是否配有内燃机？", "配有内燃机", "https://afdc.energy.gov/vehicles/electric-basics-phev"),
    ("插混车辆制动时，电机可作为发电机回收部分动能。", "插混车辆制动时电机有什么作用？", "回收部分动能", "https://afdc.energy.gov/vehicles/electric-basics-phev"),
    ("纯电动车的运动部件和需更换的液体通常少于传统燃油车。", "纯电动车的维护项目为何可能较少？", "运动部件和需更换的液体", "https://afdc.energy.gov/vehicles/electric-maintenance"),
    ("电动车事故应急处置人员应参考制造商指南和培训资料。", "电动车事故应急处置人员应参考什么？", "制造商指南和培训资料", "https://afdc.energy.gov/vehicles/electric-maintenance"),
    ("公共充电设施可扩大纯电动车日常可用行驶范围。", "公共充电设施对日常出行有什么作用？", "扩大纯电动车日常可用行驶范围", "https://afdc.energy.gov/fuels/electricity-charging-public"),
]

PROFILES = [
    {"model": "云途A1", "latin": "Yuntu A1", "year": 2025, "range": "510 公里", "seats": "5 座", "charge": "31 minutes", "warranty": "5 年或 12 万公里", "faq": "手机蓝牙钥匙", "voc": "语音唤醒偏慢", "competitor": "竞品 C1", "service": "每 12 个月"},
    {"model": "云途A1", "latin": "Yuntu A1", "year": 2026, "range": "545 公里", "seats": "5 座", "charge": "27 minutes", "warranty": "6 年或 15 万公里", "faq": "手机蓝牙钥匙", "voc": "空调风量记忆失效", "competitor": "竞品 C2", "service": "每 18 个月"},
    {"model": "云途A2", "latin": "Yuntu A2", "year": 2025, "range": "590 公里", "seats": "6 座", "charge": "34 minutes", "warranty": "5 年或 15 万公里", "faq": "远程开启空调", "voc": "第三排上下车不便", "competitor": "竞品 D1", "service": "每 12 个月"},
    {"model": "云途A2", "latin": "Yuntu A2", "year": 2026, "range": "620 公里", "seats": "6 座", "charge": "29 minutes", "warranty": "6 年或 18 万公里", "faq": "远程开启空调", "voc": "二排杯架较浅", "competitor": "竞品 D2", "service": "每 18 个月"},
]


def write_doc(stem: str, kind: str, fmt: str, text: str, title: str) -> dict:
    path = OUT / f"{stem}.{fmt}"
    if fmt == "md":
        path.write_text(f"# {title}\n{text}\n", encoding="utf-8")
        location = title
    elif fmt == "txt":
        path.write_text(text + "\n", encoding="utf-8")
        location = "正文"
    elif fmt == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["主题", "内容"])
            writer.writerow([title, text])
        location = "行 2"
    elif fmt == "docx":
        doc = Document()
        doc.add_heading(title, 1)
        doc.add_paragraph(text)
        raw = io.BytesIO()
        doc.save(raw)
        normalized = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source_zip:
            with zipfile.ZipFile(normalized, "w") as target_zip:
                for member in source_zip.infolist():
                    info = zipfile.ZipInfo(member.filename, (1980, 1, 1, 0, 0, 0))
                    info.compress_type = member.compress_type
                    info.external_attr = member.external_attr
                    target_zip.writestr(info, source_zip.read(member.filename))
        path.write_bytes(normalized.getvalue())
        location = f"章节 {title} · 段落 1"
    elif fmt == "pdf":
        page = canvas.Canvas(str(path), pagesize=A4, invariant=1)
        page.drawString(54, 780, text)
        page.save()
        location = "第 1 页"
    else:
        raise ValueError(fmt)
    return {"source": path.name, "location": location, "format": fmt}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    CASES.parent.mkdir(parents=True, exist_ok=True)
    facts = []
    manifest = []
    public_formats = ["md", "txt", "csv", "docx"] * 3
    for index, (text, question, answer, url) in enumerate(PUBLIC, 1):
        meta = write_doc(f"公开知识_{index:02d}", "public_derived", public_formats[index - 1], text, "电动车知识")
        facts.append({**meta, "question": question, "answer": answer, "required_text": answer, "kind": "public"})
        manifest.append({**meta, "origin": "public_derived", "source_url": url,
                         "access_date": "2026-10-05", "fact_mapping": text,
                         "note": "原创中文概述；不代表中国标准或车型承诺"})
    facets = [
        ("spec", "md", "产品配置"), ("charge", "pdf", "Charging"),
        ("warranty", "docx", "整车质保"), ("faq", "txt", "用户FAQ"),
        ("voc", "csv", "VOC记录"), ("competitor", "md", "竞品观察"),
        ("service", "docx", "售后周期"),
    ]
    by_profile = {}
    for profile in PROFILES:
        key = f"{profile['model']}_{profile['year']}"
        by_profile[key] = {}
        for facet, fmt, title in facets:
            prefix = f"{profile['model']} {profile['year']} 款"
            if facet == "spec":
                body = f"{prefix}的演示续航为 {profile['range']}，座位布局为 {profile['seats']}。"
                answer = profile["range"]
                question = f"{prefix}的演示续航是多少？"
            elif facet == "charge":
                body = f"{profile['latin']} {profile['year']} DC 30%-80% charging: {profile['charge']}."
                answer = profile["charge"]
                question = f"{profile['latin']} {profile['year']} DC 30%-80% charging takes how long?"
            elif facet == "warranty":
                body = f"{prefix}整车质保为 {profile['warranty']}。"
                answer = profile["warranty"]
                question = f"{prefix}整车质保多久？"
            elif facet == "faq":
                body = f"{prefix}支持{profile['faq']}功能。"
                answer = profile["faq"]
                question = f"{prefix}支持什么手机功能？"
            elif facet == "voc":
                body = f"{prefix}的合成用户反馈主题是{profile['voc']}。"
                answer = profile["voc"]
                question = f"{prefix}的 VOC 反馈主题是什么？"
            elif facet == "competitor":
                body = f"{prefix}在合成竞品观察中对照车型为{profile['competitor']}。"
                answer = profile["competitor"]
                question = f"{prefix}的竞品对照车型是哪款？"
            else:
                body = f"{prefix}建议{profile['service']}检查一次常规项目。"
                answer = profile["service"]
                question = f"{prefix}常规项目建议多久检查一次？"
            meta = write_doc(f"{key}_{facet}", "synthetic", fmt, body, title)
            fact = {**meta, "question": question, "answer": answer, "required_text": answer,
                    "kind": facet, "profile": key, "model": profile["model"], "year": profile["year"]}
            facts.append(fact)
            by_profile[key][facet] = fact
            manifest.append({**meta, "origin": "synthetic", "source_url": None, "access_date": None,
                             "fact_mapping": body, "note": "虚构车型与演示数值"})

    cases = []
    def add(group: str, question: str, selected=(), alternatives=(), answerable=True):
        evidence = [{"source": fact["source"], "location": fact["location"],
                     "required_text": fact["required_text"]} for fact in selected]
        accepted = [[fact["answer"], *(alternatives[index] if index < len(alternatives) else ())]
                    for index, fact in enumerate(selected)] if answerable else []
        cases.append({"id": f"e{len(cases)+1:03d}", "question": question, "answerable": answerable,
                      "expected_sources": [fact["source"] for fact in selected],
                      "expected_phrases": [fact["answer"] for fact in selected],
                      "acceptable_answers": accepted, "gold_evidence": evidence,
                      "group": group, "document_format": "mixed" if len({fact["format"] for fact in selected}) > 1 else
                      (selected[0]["format"] if selected else "none")})

    # Freeze the questions before any retrieval-rule tuning. No case ID is used by the app.
    for fact in facts[:12]:
        add("direct_faq", fact["question"], [fact])
    synthetic = facts[12:]
    for index in range(28):
        fact = synthetic[index]
        add("direct_faq", fact["question"], [fact])
    paraphrases = [
        "请查一下{label}的{item}，只看资料里的说法。",
        "{label}这项{item}资料怎么写？",
        "我想知道{label}关于{item}的具体记录。",
        "按现有材料，{label}的{item}是什么？",
    ]
    for index in range(30):
        fact = synthetic[(index * 7 + index // 4) % len(synthetic)]
        label = f"{fact['model']} {fact['year']} 款"
        item = {"spec": "演示续航", "charge": "30%到80%充电时间", "warranty": "整车质保",
                "faq": "手机功能", "voc": "用户反馈主题", "competitor": "对照车型",
                "service": "常规检查周期"}[fact["kind"]]
        if fact["kind"] == "charge":
            label = f"Yuntu {fact['model'][-2:]} {fact['year']}"
        add("paraphrase", paraphrases[index % 4].format(label=label, item=item), [fact])
    for index in range(35):
        fact = synthetic[(index * 5) % len(synthetic)]
        label = f"{fact['model']} {fact['year']} 款"
        if fact["kind"] == "charge":
            label = f"Yuntu {fact['model'][-2:]} {fact['year']}"
        opening = ("不要混淆其他年款或车型：" if index < 28 else "只按指定版本，不引用相近车型：")
        subject = {"spec": "演示续航", "charge": "30%到80%充电时间", "warranty": "整车质保",
                   "faq": "手机功能", "voc": "用户反馈主题", "competitor": "竞品对照车型",
                   "service": "常规检查周期"}[fact["kind"]]
        add("entity_version_confusion", f"{opening}{label} 的{subject}是什么？", [fact])
    profiles = list(by_profile.values())
    pairs = [("spec", "warranty"), ("spec", "service"), ("warranty", "faq"),
             ("faq", "voc"), ("spec", "competitor"), ("warranty", "service"), ("spec", "voc")]
    for index in range(25):
        profile = PROFILES[index % 4]
        first, second = pairs[index // 4]
        selection = [by_profile[f"{profile['model']}_{profile['year']}"][first],
                     by_profile[f"{profile['model']}_{profile['year']}"][second]]
        labels = {"spec": "演示续航", "warranty": "整车质保", "service": "常规检查周期",
                  "faq": "手机功能", "voc": "用户反馈主题", "competitor": "竞品对照车型"}
        add("cross_document", f"{profile['model']} {profile['year']} 款：{labels[first]}和{labels[second]}分别是什么？", selection)
    for index in range(20):
        fact = [item for item in synthetic if item["kind"] in {"voc", "competitor"}][index % 8]
        angle = ("具体内容", "材料记载", "结论")[(index // 8) % 3]
        subject = "用户反馈" if fact["kind"] == "voc" else "竞品对照"
        add("voc_competitor", f"针对{fact['model']} {fact['year']} 款，{subject}的{angle}是什么？", [fact])
    missing_topics = ["官方售价", "电池供应商", "车身颜色数量", "碰撞测试星级", "海外上市日期",
                      "免费换电次数", "电池详细化学配方", "年度销量", "自动驾驶级别", "零百加速时间"]
    for index in range(50):
        profile = PROFILES[index % 4]
        topic = missing_topics[index // 5]
        question = f"{profile['model']} {profile['year']} 款的{topic}是多少？"
        if index % 5 == 4:
            question = f"资料是否确认{profile['model']} {profile['year']} 款的{topic}？给出具体数值。"
        add("missing_false_premise", question, answerable=False)

    assert len(manifest) == 40 and len(cases) == 200
    assert len({item["question"] for item in cases}) == 200
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    CASES.write_text("\n".join(json.dumps(item, ensure_ascii=False) for item in cases) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
